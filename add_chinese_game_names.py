import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PATH = r"C:\Users\Admin\Documents\H5\手机远程操作适配分析.xlsx"


CHINESE_NAMES = {
    "ShengHuaWeiJi9AHQ": "生化危机9（版本需确认）",
    "JianXing": "剑星",
    "Grand Theft Auto V Enhanced": "侠盗猎车手5 增强版",
    "Romance_of_the_Three_Kingdoms_14": "三国志14",
    "Palworld_V1.0.0.100427": "幻兽帕鲁",
    "DiPingXian6": "地平线6（需确认）",
    "DuShiTianJiXian2": "城市：天际线2",
    "BingQiShiDai2": "冰汽时代2",
    "ShengHuaWeiJi4CZB": "生化危机4 重制版",
    "ZSGWSQiYuan": "真三国无双：起源",
    "ShenHaiMiHang2YXSY": "深海迷航2（版本需确认）",
    "PalworldV7.1.86065": "幻兽帕鲁",
    "NvShenYiWenLu5HJB": "女神异闻录5 皇家版",
    "Thief Simulator 2": "小偷模拟器2",
    "Grand.Theft.Auto.The.Definitive.Edition": "侠盗猎车手：三部曲 最终版",
    "Victoria3": "维多利亚3",
    "Monster Hunter World Iceborne": "怪物猎人：世界 冰原",
    "ShiZhiCunZai": "逝之存在（需确认）",
    "ELDEN RING DLC": "艾尔登法环 DLC",
    "Battlefield 6": "战地6",
    "FuShe4": "辐射4",
    "QuanMianZhanZhengZC": "全面战争：战锤（需确认）",
    "Just Cause 3": "正当防卫3",
    "ShaLuJianTa2": "杀戮尖塔2",
    "XiaoShiDeGuangMangKS": "消逝的光芒（版本需确认）",
    "Kingdom Come Deliverance II": "天国：拯救2",
    "RenWang3": "仁王3",
    "FangZhouFangKuaiShiJie": "方舟：方块世界",
    "JingCaiMoNiQiXunJing": "警察模拟器：巡警",
    "ShuangYingQiJing": "双影奇境",
    "AnHeiPoHuaiShen2YHCSCZLYB": "暗黑破坏神2：狱火重生（版本需确认）",
    "ChangTuLvXing": "长途旅行",
    "Left.4.Dead.2": "求生之路2",
    "NARUTO.X.BORUTO.Ultimate.Ninja.STORM.CONNECTIONS": "火影忍者X博人传：终极忍者风暴羁绊",
    "A081": "A081（需确认）",
    "ShiMingZhaoHuan6XDZZCZB": "使命召唤6：现代战争2 重制版",
    "Hitman 3": "杀手3",
    "WuShuangShenYuan": "无双深渊",
    "HeiAnZhiHun3": "黑暗之魂3",
    "CiKeXinTiaoHJ": "刺客信条：幻景",
    "ZangYuanLingZhu": "葬渊领主（需确认）",
    "ShanGuoZhi8CZBWLJQB": "三国志8 重制版 威力加强版（需确认）",
    "RenZhongZhiLong0": "人中之龙0",
    "Ghost of Tsushima DIRECTOR'S CUT": "对马岛之魂 导演剪辑版",
    "QuanHuang15": "拳皇15",
    "XiaoXiaoMengYan3": "小小梦魇3",
    "BoBiDeYouXiShiGuang": "波比的游戏时间",
    "QiLuLvRen0": "歧路旅人0（需确认）",
    "KongZhiGuiJiTHE1ST": "空之轨迹 the 1st",
    "ZhiWuDaZhanJiangShiWuShuangBan0.1": "植物大战僵尸无双版",
    "Warhammer 40000 Space Marine 2": "战锤40000：星际战士2",
    "LieRenHuangYeDeZhaoHuan": "猎人：荒野的召唤",
    "JiYuan1800": "纪元1800",
    "Darkest Dungeon 2V2.00.75033": "暗黑地牢2",
    "Age.of.Empires.II.Definitive.Edition.Victors.and.Vanquished-P2P": "帝国时代2：决定版 胜者与败者",
    "DiGuoShiDai4": "帝国时代4",
    "Skul": "小骨：英雄杀手",
    "JiJingLingF": "寂静岭f",
    "HuoGeWoZiZhiYi": "霍格沃茨之遗",
    "Assassin": "刺客信条（具体版本需确认）",
    "Far Cry Primal": "孤岛惊魂：原始杀戮",
    "DX": "DX（需确认）",
    "QuanMianZhanZhengSG": "全面战争：三国",
    "XueRanXiaoZhen": "血染小镇",
    "ShengHuaWeiJi8": "生化危机8：村庄",
    "QunXingYinHeBan": "群星 银河版",
    "HaiDiDaLieSha": "海底大猎杀",
    "LangRenJueQi": "浪人崛起",
    "Craftopia.Build.13912084": "创世理想乡",
    "BLEACHHunPoJueXing": "死神：魂魄觉醒",
    "TaoSheng2": "逃生2",
    "SILENT HILL 2": "寂静岭2",
    "HuangYeDaBiaoKe": "荒野大镖客",
    "PiNuoCaoDeHuangYan": "匹诺曹的谎言",
    "ChengShiTianJiXian": "城市：天际线",
    "YeGouZiLieTouGuai": "野狗子：裂头怪",
    "Contra Operation Galuga": "魂斗罗：加鲁加行动",
    "ShuangJieLongZL": "双截龙：再临",
    "RenZheLongJianZhuan4": "忍者龙剑传4",
    "QuanMianChongTuDiKang": "全面冲突：抵抗",
    "SuDanDeYouXi": "苏丹的游戏",
    "DRAGON BALL Sparking ZERO Hero of Justice": "七龙珠 电光炸裂！ZERO 正义英雄",
    "FaHuanHeiYeJunLin": "艾尔登法环：黑夜君临",
    "TEKKEN 8": "铁拳8",
    "ShiMingZhaoHuan9HSXD2": "使命召唤9：黑色行动2",
    "Hollow Knight Silksong": "空洞骑士：丝之歌",
    "OuLuFengYun5": "欧陆风云5",
    "Only Up": "只有向上",
    "ShiJianLvZheCSSG": "时间旅者（需确认）",
    "ChaoJiHeiAnQiPian": "超级黑暗欺骗",
    "EMoLunPan": "恶魔轮盘",
    "ShanDian11Ren": "闪电十一人",
    "ShiLuoZhiHun": "失落之魂",
    "ChaoShiMoNiQi": "超市模拟器",
    "ShengLingChongSuo": "圣灵重塑（需确认）",
    "QiGaiMoNiQi": "乞丐模拟器",
    "WMJSJX": "WMJSJX（需确认）",
    "KAGE Shadow of the Ninja": "影之忍者",
    "A2084": "A2084（需确认）",
    "Dying Light 2 DLC": "消逝的光芒2 DLC",
    "QiYiRenShengShuangChongBaoGuang": "奇异人生：双重曝光",
    "VRZhanShi5": "VR战士5",
    "Expedition 33": "光与影：33号远征",
    "WuSuoWangGuo": "雾锁王国",
    "ShengJianChuanShuo": "圣剑传说",
    "Risk.of.Rain.2.v1.10a": "雨中冒险2",
    "MuChangWuYu": "牧场物语",
    "NoRestForTheWicked": "恶意不息",
    "Ghostrunner 2": "幽灵行者2",
    "A490": "A490（需确认）",
    "QiNvShenZhiDao": "祇：女神之道",
    "SaiErDaChuanShuoKuangYeZhiXi": "塞尔达传说：旷野之息",
    "XiaoMaoMiDaChengShi": "小猫咪大城市",
    "Backpack.Hero": "背包英雄",
    "QiLongZhuDianGuangZhaLie": "七龙珠 电光炸裂",
    "LongZhiXinTiao2": "龙之信条2",
    "RenZheLongJianZhuan2": "忍者龙剑传2",
    "Dragon Ball Z Kakarot 23rd World Tournament": "龙珠Z：卡卡罗特 第23届天下第一武道会",
    "The.Rogue.Prince.of.Persia": "波斯王子：Rogue",
    "PATAPON12": "啪嗒砰1+2",
    "MengHuanMoNIZhan1and2": "梦幻模拟战1&2",
    "DRAGON QUEST TREASURES": "勇者斗恶龙 寻宝探险团",
    "FeiYeChuanQiCZB": "菲耶传奇 重制版（需确认）",
    "yuzu-windows-msvc-early-access-v1.2.1": "Yuzu模拟器 早期访问版",
    "Sid Meiers Civilization VI": "席德·梅尔的文明6",
    "CRISIS.CORE.FINAL.FANTASY.VII.REUNION.DIGITAL.DELUXE.EDITION.v1.0.3": "最终幻想7 核心危机 Reunion 数字豪华版",
    "LongTengShiJiYZSHZ": "龙腾世纪：影障守护者",
    "JiangShiHuiMieGongCheng": "僵尸毁灭工程",
    "Shadow Tactics Blades of the Shogun": "影子战术：将军之刃",
    "JiJiaZhanMoSHZY": "机甲战魔：神话之裔",
    "MoNiRenSheng4": "模拟人生4",
    "NueShaRongLu": "虐杀熔炉（需确认）",
    "LianZaiYiQi": "连在一起",
    "YongChuangSiRenGuAHZR": "勇闯死人谷：暗黑之日",
    "RenZheShenGui": "忍者神龟",
    "Granblue.Fantasy.Relink": "碧蓝幻想 Relink",
    "Another.Crabs.Treasure-GoldBerg": "蟹蟹寻宝奇遇",
    "QianShuiYuanDaiFu": "潜水员戴夫",
    "Avowed": "宣誓",
    "8HaoZhanTai": "8号站台",
    "XieEMingKeKaiXiMoZhu": "邪恶冥刻：凯西模组",
    "Atlas.Fallen": "尘封大陆",
    "Lords of the Fallen": "堕落之主",
    "Fate.Samurai.Remnant": "Fate/Samurai Remnant",
    "MiYingManDeLa": "迷影曼德拉（需确认）",
    "Shovel.Knight.Pocket.Dungeon.Puzzlers.Pack.v2.0.3": "铲子骑士：口袋地牢 解谜包",
    "TCG.Card.Shop.Simulator.v0.32-105709": "TCG卡牌商店模拟器",
    "Orcs.Must.Die.3.v1.2.5.1-P2P": "兽人必须死3",
    "BangOnBalls": "砰砰球历险记",
    "Risk.of.Rain.Returns": "雨中冒险 回归",
    "SWORD ART ONLINE Fractured Daydream": "刀剑神域 碎梦边境",
    "TaiTanZhiLv2": "泰坦之旅2",
    "JiaMianQiShiDianFengLuanDou": "假面骑士 巅峰乱斗",
    "TaoSheng1GaoMiZhe": "逃生：告密者",
    "Spell.Disk.Build": "法术圆盘",
    "XiaoChouPai": "小丑牌",
    "ManNiHanBaoDian": "曼尼汉堡店（需确认）",
    "Bloody Heaven 2": "血腥天堂2",
    "Sea of Stars": "星之海",
    "YouMiYaDeLianJinGongFang": "优米雅的炼金工房",
    "Hades.II.v0.89923a": "哈迪斯2",
    "LaBiXiaoXinWoYuBoShiDeShuJia": "蜡笔小新：我与博士的暑假",
    "ManWeiVSKaPuKong": "漫威VS卡普空",
    "LeGaoDaLuanDou": "乐高大乱斗",
    "LaLiSaiDeYiShu": "拉莉萨的艺术（需确认）",
    "SiShenVSHuoYing Flight": "死神VS火影 Flight",
    "FeiChangPuTongDeLu": "非常普通的鹿",
    "JiXingDongWuYuan": "奇兴动物园（需确认）",
    "BaoKeMengWuXianRongHe": "宝可梦无限融合",
    "JueDiQiuSheng2": "绝地求生2",
}


def has_cjk(text):
    return bool(re.search(r"[\u4e00-\u9fff]", text))


def chinese_name(original):
    if original in CHINESE_NAMES:
        return CHINESE_NAMES[original]
    if has_cjk(original):
        return original
    return f"{original}（需确认）"


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb["手机远程操作适配分析"]

    headers = [cell.value for cell in ws[1]]
    if "中文游戏名" not in headers:
        ws.insert_cols(3)
        ws.cell(row=1, column=2).value = "原始游戏名"
        ws.cell(row=1, column=3).value = "中文游戏名"
    else:
        ws.cell(row=1, column=2).value = "原始游戏名"

    for row in range(2, ws.max_row + 1):
        original = str(ws.cell(row=row, column=2).value or "")
        ws.cell(row=row, column=3).value = chinese_name(original)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [8, 36, 38, 12, 62, 34]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(PATH)
    print(PATH)


if __name__ == "__main__":
    main()
