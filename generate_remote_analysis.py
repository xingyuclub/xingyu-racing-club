import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SRC = r"C:\Users\Admin\Desktop\6.1-7.27热门游戏.xlsx"
OUT = r"C:\Users\Admin\Documents\H5\手机远程操作适配分析.xlsx"


KNOWN = {
    "Romance_of_the_Three_Kingdoms_14": ("适合", "回合/策略经营节奏慢，主要是菜单和规划操作，远程延迟影响小。", "触控即可，建议横屏"),
    "Victoria3": ("适合", "大战略/经营类，暂停和菜单操作多，不依赖即时反应。", "触控或蓝牙鼠标更好"),
    "JiYuan1800": ("适合", "城市建设/经营管理为主，可暂停规划，手机远程可控。", "触控或蓝牙鼠标更好"),
    "Sid Meiers Civilization VI": ("适合", "回合制策略，不吃操作速度，适合碎片时间远程操作。", "触控即可"),
    "Shadow Tactics Blades of the Shogun": ("适合", "战术潜入可暂停/规划路线，反应压力相对低。", "触控可玩，鼠标更舒服"),
    "博德之门3": ("适合", "回合制战斗和剧情选择为主，延迟影响很小。", "触控可玩，手柄更舒服"),
    "永恒之柱2": ("适合", "暂停制/战术RPG，可慢慢下指令，适合远程。", "触控可玩，建议横屏"),
    "春逝百年抄": ("适合", "互动影像/推理剧情为主，操作少，手机远程体验好。", "触控即可"),
    "TCG.Card.Shop.Simulator.v0.32-105709": ("适合", "经营模拟节奏慢，主要是点击、整理和管理，远程友好。", "触控即可"),
    "Backpack.Hero": ("适合", "背包构筑/回合制玩法，主要是拖拽和决策，延迟不敏感。", "触控即可"),
    "Shovel.Knight.Pocket.Dungeon.Puzzlers.Pack.v2.0.3": ("适合", "解谜/消除向，节奏可控，短局适合手机远程。", "触控即可"),
    "Sea of Stars": ("适合", "回合制RPG，操作简单，低延迟要求。", "触控或手柄"),
    "DRAGON QUEST TREASURES": ("勉强适合", "偏轻动作RPG，探索可远程玩，但战斗和镜头操作会受延迟影响。", "建议蓝牙手柄"),
    "哆啦A梦牧场物语": ("适合", "牧场经营和采集为主，节奏慢，适合远程慢玩。", "触控或手柄"),
    "MuChangWuYu": ("适合", "牧场/生活模拟倾向，节奏慢，远程操作压力低。", "触控或手柄"),
    "MoNiRenSheng4": ("适合", "生活模拟/建造管理为主，鼠标点击多，反应要求低。", "触控或蓝牙鼠标更好"),
    "ChaoShiMoNiQi": ("适合", "模拟经营类，主要是管理和摆放，低延迟可接受。", "触控即可"),
    "QiGaiMoNiQi": ("适合", "模拟生存/管理节奏相对可控，不是强竞技操作。", "触控或手柄"),
    "QiYiRenShengShuangChongBaoGuang": ("适合", "剧情选择/叙事推理倾向，操作轻，远程友好。", "触控即可"),
    "LaBiXiaoXinWoYuBoShiDeShuJia": ("适合", "休闲剧情/生活冒险，节奏慢，适合手机远程。", "触控或手柄"),
    "YouMiYaDeLianJinGongFang": ("适合", "炼金工房类偏采集、养成、菜单管理，延迟影响不大。", "触控或手柄"),
    "XiaoChouPai": ("适合", "卡牌/策略构筑倾向，主要是选择和出牌，适合远程。", "触控即可"),
    "ManNiHanBaoDian": ("适合", "烹饪/模拟经营倾向，操作节奏通常可控，适合远程轻玩。", "触控即可"),
    "BaoKeMengWuXianRongHe": ("适合", "宝可梦融合/回合制倾向，菜单和回合决策为主。", "触控即可"),
    "QunXingYinHeBan": ("适合", "群星类大战略，暂停规划和菜单操作多，适合远程。", "触控或蓝牙鼠标更好"),
    "OuLuFengYun5": ("适合", "欧陆风云类大战略，暂停和菜单规划为主，低反应需求。", "触控或蓝牙鼠标更好"),
    "ChengShiTianJiXian": ("适合", "城市天际线类城市建设，规划和菜单操作为主。", "触控或蓝牙鼠标更好"),
    "DuShiTianJiXian2": ("适合", "城市建设/经营管理，节奏慢，可暂停，适合远程操作。", "触控或蓝牙鼠标更好"),
    "BingQiShiDai2": ("适合", "冰汽时代类城市生存策略，决策和规划为主，延迟影响小。", "触控或蓝牙鼠标更好"),
    "JingCaiMoNiQiXunJing": ("适合", "模拟经营/巡警类倾向，实时压力有限，远程可玩。", "触控或手柄"),
    "JiXingDongWuYuan": ("适合", "动物园/经营模拟倾向，规划和管理为主，适合远程。", "触控或蓝牙鼠标"),
    "8HaoZhanTai": ("适合", "找异常/步行模拟类，操作简单，反应要求低。", "触控即可"),
    "迷失": ("适合", "探索解谜为主，动作压力较低，手机远程体验较好。", "触控或手柄"),
    "蔑视": ("勉强适合", "第一人称解谜探索可远程，但视角和战斗段落会受延迟影响。", "建议手柄"),
    "微观世界大冒险": ("适合", "休闲冒险/解谜倾向，操作量较少，适合手机远程。", "触控或手柄"),
    "求生岛：不老泉传说": ("勉强适合", "生存探索可玩，但采集、战斗和镜头控制在手机上偏累。", "建议手柄"),
    "死亡搁浅": ("勉强适合", "探索运输节奏不算快，但长时间移动和镜头控制更适合手柄。", "建议手柄"),
    "无人深空": ("勉强适合", "探索建造可以远程玩，但飞行和战斗对触控不友好。", "建议手柄"),
    "星空": ("勉强适合", "RPG探索和对话可远程，射击和飞船战斗受延迟影响。", "建议手柄"),
}

UNSUITABLE_KEYWORDS = [
    "使命召唤", "ShiMingZhaoHuan", "Battlefield", "光环", "Left.4.Dead", "TaoSheng", "JueDiQiuSheng",
    "Far Cry", "孤岛惊魂", "无主之地", "QuanHuang", "TEKKEN", "STORM", "ManWeiVSKaPuKong",
    "Dragon Ball", "DRAGON BALL Sparking", "QiLongZhuDianGuangZhaLie", "VRZhanShi",
    "极品飞车", "地平线5", "DiPingXian6", "Ghostrunner", "只狼", "ELDEN", "艾尔登",
    "HeiAnZhiHun", "Lords of the Fallen", "NoRestForTheWicked", "RenWang",
    "Hollow Knight", "空洞骑士", "茶杯头", "Only Up", "Contra", "KAGE", "忍者神龟",
    "Risk.of.Rain", "哈迪斯", "Hades", "恶魔五月哭", "鬼泣", "Monster Hunter",
    "盗贼之海", "Warhammer 40000 Space Marine 2", "消逝的光芒", "Dying Light", "森林之子",
    "生化危机", "ShengHuaWeiJi", "SILENT HILL", "地铁离去", "原子之心", "虐杀原形", "NueSha",
    "Granblue.Fantasy.Relink", "破晓传说", "Atlas.Fallen", "Fate.Samurai.Remnant", "Avowed",
]

FIT_KEYWORDS = [
    "Three_Kingdoms", "Victoria", "Civilization", "QunXing", "OuLu", "JiYuan",
    "MoNi", "Simulator", "模拟", "牧场", "MuChang", "TianJiXian", "城市", "Backpack", "Card",
    "TCG", "Puzzlers", "Sea of Stars", "永恒之柱", "博德", "Romance", "Shadow Tactics",
    "Darkest Dungeon", "DRAGON QUEST", "炼金工房", "LianJin", "BaoKeMeng", "XiaoChouPai",
    "8HaoZhanTai", "春逝百年抄", "LaBiXiaoXin",
]

MAYBE_KEYWORDS = [
    "Grand Theft Auto", "荒野大镖客", "骑马与砍杀", "刺客信条", "Assassin", "古墓丽影",
    "战神", "Just Cause", "正当防卫", "漫威蜘蛛侠", "看门狗", "Hitman", "龙珠Z", "NARUTO",
    "海贼无双", "尼尔", "FINAL FANTASY", "最终幻想", "奥日", "Skul", "Palworld", "Craftopia",
    "HuoGeWoZiZhiYi", "Kingdom Come", "审判之眼", "热血少女", "Orcs.Must.Die", "渡神纪",
    "SWORD ART ONLINE", "CRISIS.CORE", "黑道圣徒", "Ghost of Tsushima", "赛博朋克", "FuShe4",
    "Another.Crabs.Treasure", "Rogue.Prince", "PATAPON", "Spell.Disk", "诸神黄昏幸存者",
    "Bloody Heaven",
]

UNKNOWN_HINTS = ["A081", "A2084", "A490", "DX", "WMJSJX"]


def classify(name):
    if name in KNOWN:
        return KNOWN[name]
    if any(keyword in name for keyword in UNKNOWN_HINTS):
        return ("需确认", "名称像内部编号或缩写，无法仅凭名称判断玩法；建议先确认对应游戏。", "先人工核对")
    if any(keyword in name for keyword in UNSUITABLE_KEYWORDS):
        return ("不适合", "偏射击、格斗、竞速、魂类或高强度动作，手机远程的延迟和触控按键会明显影响体验。", "不建议；必须玩则用手柄并保证低延迟网络")
    if any(keyword in name for keyword in FIT_KEYWORDS):
        return ("适合", "偏策略、模拟经营、卡牌、回合制或解谜，主要靠选择和规划，远程延迟影响小。", "触控即可，复杂菜单建议横屏/蓝牙鼠标")
    if any(keyword in name for keyword in MAYBE_KEYWORDS):
        return ("勉强适合", "偏动作冒险或开放世界，探索和剧情可远程，战斗、驾驶或平台跳跃会受延迟影响。", "建议蓝牙手柄，避免高难度战斗")
    return ("需确认", "仅凭当前名称难以稳定识别玩法类型；不建议直接作为适合远程的候选。", "先确认具体游戏类型")


def style_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main():
    source_wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    source_ws = source_wb[source_wb.sheetnames[0]]
    games = [row[0] for row in source_ws.iter_rows(min_row=2, values_only=True) if row and row[0]]

    rows = []
    for index, game in enumerate(games, 1):
        level, reason, suggestion = classify(str(game))
        rows.append([index, game, level, reason, suggestion])

    out_wb = Workbook()
    detail_ws = out_wb.active
    detail_ws.title = "手机远程操作适配分析"
    detail_ws.append(["序号", "游戏名", "适合程度", "理由", "操作建议"])
    for row in rows:
        detail_ws.append(row)

    level_fills = {
        "适合": PatternFill("solid", fgColor="C6EFCE"),
        "勉强适合": PatternFill("solid", fgColor="FFEB9C"),
        "不适合": PatternFill("solid", fgColor="F4CCCC"),
        "需确认": PatternFill("solid", fgColor="D9EAF7"),
    }
    style_sheet(detail_ws)
    for row in detail_ws.iter_rows(min_row=2):
        row[2].fill = level_fills.get(row[2].value, PatternFill())

    for index, width in enumerate([8, 42, 12, 62, 34], 1):
        detail_ws.column_dimensions[get_column_letter(index)].width = width
    detail_ws.freeze_panes = "A2"
    detail_ws.auto_filter.ref = detail_ws.dimensions

    summary_ws = out_wb.create_sheet("汇总")
    summary_ws.append(["适合程度", "数量", "说明"])
    counts = {level: sum(1 for row in rows if row[2] == level) for level in ["适合", "勉强适合", "不适合", "需确认"]}
    notes = {
        "适合": "优先推荐手机远程：低反应需求、菜单/回合/经营/解谜为主。",
        "勉强适合": "可玩但建议手柄：探索剧情可接受，高强度操作体验会下降。",
        "不适合": "不优先推荐：对延迟、视角、按键或精准操作要求高。",
        "需确认": "名称不明确或缩写较多，需要人工确认具体玩法后再判断。",
    }
    for level in ["适合", "勉强适合", "不适合", "需确认"]:
        summary_ws.append([level, counts[level], notes[level]])
    style_sheet(summary_ws)
    for index, width in enumerate([14, 10, 70], 1):
        summary_ws.column_dimensions[get_column_letter(index)].width = width

    out_wb.save(OUT)
    print(OUT)
    print(counts)


if __name__ == "__main__":
    main()
